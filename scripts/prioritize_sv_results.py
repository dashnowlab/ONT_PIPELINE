#!/usr/bin/env python3

import sys
import os
import pandas as pd
import numpy as np
import re
from openpyxl import load_workbook


def count_genes_rowwise(val):
    if pd.isna(val) or str(val).strip() == "":
        return np.nan
    return len([x for x in str(val).split(",") if x.strip()])


def any_gene_in_acmg_comma(genes_str, acmg_set):
    if pd.isna(genes_str) or str(genes_str).strip() == "":
        return False
    genes = [x.strip() for x in str(genes_str).split(",") if x.strip()]
    return any(g in acmg_set for g in genes)


def normalize_sidra_lims_id(x):
    if pd.isna(x):
        return x
    x = re.sub(r"\s+", "", str(x))
    if re.search(r"^(200|201|210|101)", x):
        x = "X0" + x
    if re.search(r"^(020)", x):
        x = "X" + x
    return x


def safe_numeric(series):
    return pd.to_numeric(series.replace("", -1), errors="coerce").fillna(-1)


def write_excel_calls_sheet(template_xlsx, output_xlsx, df):
    wb = load_workbook(template_xlsx)

    if "calls" in wb.sheetnames:
        ws = wb["calls"]
        wb.remove(ws)

    ws = wb.create_sheet("calls")

    for col_idx, col_name in enumerate(df.columns, start=1):
        ws.cell(row=1, column=col_idx, value=col_name)

    for row_idx, row in enumerate(df.itertuples(index=False), start=2):
        for col_idx, value in enumerate(row, start=1):
            if pd.isna(value):
                value = None
            ws.cell(row=row_idx, column=col_idx, value=value)

    wb.save(output_xlsx)


def main():
    if len(sys.argv) != 9:
        sys.stderr.write(
            "Usage:\n"
            "python prioritize_sv_results.py "
            "<sv_file> <sample_ids_xlsx_or_empty> <acmg_gene_list_or_empty> "
            "<svtk_bed> <genotypes_tsv> <output_filtered> <output_acmg> <output_excel>\n"
        )
        sys.exit(1)

    sv_file_path    = sys.argv[1]
    sample_ids_path = sys.argv[2]
    acmg_gene_path  = sys.argv[3]
    svtk_bed_path   = sys.argv[4]
    genotypes_path  = sys.argv[5]
    output_filtered = sys.argv[6]
    output_acmg     = sys.argv[7]
    output_excel    = sys.argv[8]

    sv_file = pd.read_csv(sv_file_path, sep="\t", dtype=str)
    svtk = pd.read_csv(svtk_bed_path, sep="\t", dtype=str)
    genotypes = pd.read_csv(genotypes_path, sep="\t", dtype=str)

    drop_cols = [
        "X.chrom", "start", "end", "svtype", "samples", "ALGORITHMS",
        "CHR2", "END", "STRANDS", "SVLEN", "SVTYPE"
    ]
    svtk = svtk.drop(columns=[c for c in drop_cols if c in svtk.columns], errors="ignore")

    df_full = sv_file.merge(genotypes, on="ID", how="left")
    df = df_full.merge(svtk, left_on="ID", right_on="name", how="left")

    df["HET"] = (df == "0|1").sum(axis=1)
    df["HOM"] = (df == "1|1").sum(axis=1)

    if "Novel" not in df.columns:
        df["Novel"] = np.nan

    max_af_num = pd.to_numeric(df["Max_AF"], errors="coerce").fillna(-1) if "Max_AF" in df.columns else pd.Series(-1, index=df.index)

    for svtype in ["DEL", "DUP", "INV", "INS"]:
        mask = (df["SV_type"] == svtype) & (max_af_num == 0)
        df.loc[mask, "Novel"] = "TRUE"

    df["Max_AF"]       = safe_numeric(df["Max_AF"])       if "Max_AF" in df.columns else -1
    df["B_loss_AFmax"] = safe_numeric(df["B_loss_AFmax"]) if "B_loss_AFmax" in df.columns else -1
    df["B_gain_AFmax"] = safe_numeric(df["B_gain_AFmax"]) if "B_gain_AFmax" in df.columns else -1
    df["B_ins_AFmax"]  = safe_numeric(df["B_ins_AFmax"])  if "B_ins_AFmax" in df.columns else -1
    df["B_inv_AFmax"]  = safe_numeric(df["B_inv_AFmax"])  if "B_inv_AFmax" in df.columns else -1
    df["SV_length"]    = pd.to_numeric(df["SV_length"], errors="coerce").fillna(0) if "SV_length" in df.columns else 0

    df_filtered = df[(df["HET"] < 5) & (df["HOM"] < 5)].copy()

    filt = (
        ((df_filtered["SV_type"] == "DEL") &
         ((df_filtered["Max_AF"] < 0.001) | (df_filtered["B_loss_AFmax"] < 0.001)) &
         (df_filtered["SV_length"] > -10000000))
        |
        ((df_filtered["SV_type"] == "DUP") &
         ((df_filtered["Max_AF"] < 0.001) | (df_filtered["B_gain_AFmax"] < 0.001)) &
         (df_filtered["SV_length"] < 10000000))
        |
        ((df_filtered["SV_type"] == "INS") &
         ((df_filtered["Max_AF"] < 0.001) | (df_filtered["B_ins_AFmax"] < 0.001)))
        |
        ((df_filtered["SV_type"] == "INV") &
         ((df_filtered["Max_AF"] < 0.001) | (df_filtered["B_inv_AFmax"] < 0.001)) &
         (df_filtered["SV_length"] < 10000000))
    )

    df_filtered = df_filtered[filt].copy()

    for col in ["REF", "ALT"]:
        if col in df_filtered.columns:
             df_filtered.drop(columns=col, inplace=True)

    selected_columns = ["PREDICTED_LOF", "PREDICTED_COPY_GAIN", "PREDICTED_INTRAGENIC_EXON_DUP"]
    for col in selected_columns:
        if col in df_filtered.columns:
            df_filtered[f"{col}_count"] = df_filtered[col].apply(count_genes_rowwise)

    if sample_ids_path and sample_ids_path != "" and os.path.exists(sample_ids_path):
        print(f"Sample ID Excel found: {sample_ids_path}", file=sys.stderr)

        sample_df = pd.read_excel(sample_ids_path, sheet_name="ids_new_with_affection_status_f")
        sample_df["Sidra LIMS ID"] = sample_df["Sidra LIMS ID"].apply(normalize_sidra_lims_id)

        original_ids = sample_df[sample_df["Sidra LIMS ID"].notna()].copy()

        rename_map = dict(
            zip(
                original_ids["Sidra LIMS ID"].astype(str),
                original_ids["Original Sample Name"].astype(str)
            )
        )

        df_filtered = df_filtered.rename(columns=rename_map)
        write_excel_calls_sheet(sample_ids_path, output_excel, df_filtered)
    else:
        print("Sample ID Excel not found. Skipping sample renaming and Excel update.", file=sys.stderr)

    df_filtered.to_csv(output_filtered, sep="\t", index=False)

    if acmg_gene_path and acmg_gene_path != "" and os.path.exists(acmg_gene_path):
        print(f"ACMG gene list found: {acmg_gene_path}", file=sys.stderr)

        acmg_gene = pd.read_csv(acmg_gene_path, sep="\t", dtype=str)
        acmg_set = set(acmg_gene["Gene"].dropna().astype(str).unique())

        if "PREDICTED_LOF" in df_filtered.columns:
            acmg_rare_plof = df_filtered[
                df_filtered["PREDICTED_LOF"].apply(lambda x: any_gene_in_acmg_comma(x, acmg_set))
            ].copy()
        else:
            acmg_rare_plof = df_filtered.iloc[0:0].copy()

        acmg_rare_plof.to_csv(output_acmg, sep="\t", index=False)
    else:
        print("ACMG gene list not found. Skipping ACMG prioritization.", file=sys.stderr)


if __name__ == "__main__":
    main()